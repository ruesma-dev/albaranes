# evals/revision/lectura.py
"""Del `.xlsx` de la revisión manual a `list[FilaPlana]`. Infraestructura.

Dos cuidados que ya han costado caro:

1. **Las columnas se localizan por NOMBRE**, nunca por posición. El Excel es
   del humano: añade columnas, las mueve, y arrastra una errata en el
   encabezado (`codigo alabran`). Leer por posición convierte cualquiera de
   las tres cosas en un reparto silenciosamente equivocado.
2. **Se lee la copia de `evals/inputs/fuente/`**, no el original de OneDrive:
   el humano lo tiene abierto a menudo y abrirlo da `PermissionError`.
"""

from __future__ import annotations

import datetime as dt
import re
import unicodedata
from pathlib import Path

from evals.revision.modelos import COLUMNAS, COLUMNAS_OBLIGATORIAS, FilaPlana

#: Copia estable del Excel de trabajo del humano. Fuera de git: lleva precios.
RUTA_FUENTE = (
    Path(__file__).resolve().parents[1] / "inputs" / "fuente" / "evals_summary.xlsx"
)

_PARENTESIS = re.compile(r"\([^)]*\)")
_NO_ALFANUMERICO = re.compile(r"[^0-9a-z]+")


class ErrorLectura(RuntimeError):
    """El Excel no está, o no trae las columnas sin las que no se reparte."""


def normalizar_encabezado(texto: object | None) -> str:
    """«Tipo de albaran (hormigon, residuos…)» → `tipo de albaran`.

    Lo que va entre paréntesis es ayuda para quien rellena la tabla, no parte
    del nombre de la columna; mismo criterio que `evals/conversor.py`.
    """
    if texto is None:
        return ""
    sin_ayuda = _PARENTESIS.sub(" ", str(texto))
    descompuesto = unicodedata.normalize("NFKD", sin_ayuda.lower())
    sin_acentos = "".join(c for c in descompuesto if not unicodedata.combining(c))
    return _NO_ALFANUMERICO.sub(" ", sin_acentos).strip()


def localizar_columnas(encabezados: list[object | None]) -> dict[str, int]:
    """Clave interna -> índice de columna. Lo que sobra se ignora."""
    normalizados = [normalizar_encabezado(e) for e in encabezados]
    indices: dict[str, int] = {}
    for clave, admitidos in COLUMNAS.items():
        for admitido in admitidos:
            if admitido in normalizados:
                indices[clave] = normalizados.index(admitido)
                break
    ausentes = [c for c in COLUMNAS_OBLIGATORIAS if c not in indices]
    if ausentes:
        raise ErrorLectura(
            f"a la tabla plana le faltan columnas obligatorias: "
            f"{', '.join(ausentes)}. Encabezados encontrados: "
            f"{', '.join(n for n in normalizados if n)}"
        )
    return indices


def _celda(valor: object | None) -> object | None:
    """Fechas a AAAA-MM-DD; el resto tal cual, sin tocar ceros a la izquierda."""
    if isinstance(valor, dt.datetime):
        return valor.date().isoformat()
    if isinstance(valor, dt.date):
        return valor.isoformat()
    if isinstance(valor, str):
        return valor.strip()
    return valor


def leer(origen: Path | str = RUTA_FUENTE, pestana: str | None = None) -> list[FilaPlana]:
    """Lee la tabla plana entera. Las filas en blanco no son filas."""
    import openpyxl  # import perezoso: solo la infraestructura necesita openpyxl

    ruta = Path(origen)
    if not ruta.is_file():
        raise ErrorLectura(
            f"no existe el fichero de origen '{ruta}'. La revisión manual vive "
            f"en OneDrive; en el repositorio solo entra la copia de "
            f"evals/inputs/fuente/, que no se versiona."
        )

    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        hoja = libro[pestana] if pestana else libro[libro.sheetnames[0]]
        filas: list[FilaPlana] = []
        indices: dict[str, int] | None = None
        for numero, celdas in enumerate(hoja.iter_rows(values_only=True), start=1):
            if indices is None:
                indices = localizar_columnas(list(celdas))
                continue
            if all(_esta_vacia(celda) for celda in celdas):
                continue
            valores = {
                clave: _celda(celdas[indice]) if indice < len(celdas) else None
                for clave, indice in indices.items()
            }
            filas.append(FilaPlana(numero_fila=numero, valores=valores))
    finally:
        libro.close()

    if indices is None:
        raise ErrorLectura(f"'{ruta}' no tiene ni fila de encabezados.")
    return filas


def _esta_vacia(valor: object | None) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())
