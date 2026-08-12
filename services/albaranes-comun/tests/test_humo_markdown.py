# tests/test_humo_markdown.py
"""Test de humo del conversor markitdown (DOCX/PDF → Markdown).

Genera un DOCX con tabla de tarifas y comprueba que la conversión
preserva la tabla como tabla Markdown. Valida también la combinación de
varias fuentes y la degradación elegante (fallback / nunca lanza).

Requiere los extras: pip install -e ".[markdown,dev]"
Si markitdown/python-docx no están, los tests dependientes se saltan.

Ejecutar:  pytest tests/test_humo_markdown.py -v
"""
from __future__ import annotations

import io

import pytest

from ruesma_comun.markdown import (
    FuenteDocumento,
    a_markdown,
    combinar_a_markdown,
    markitdown_disponible,
)

requiere_markitdown = pytest.mark.skipif(
    not markitdown_disponible(), reason="markitdown no instalado"
)


def _docx_con_tabla() -> bytes | None:
    try:
        from docx import Document
    except Exception:
        return None
    doc = Document()
    doc.add_heading("CONTRATO 0464", level=1)
    doc.add_paragraph("Suministro de hormigón para la obra 0464.")
    filas = [
        ("Concepto", "Unidad", "Precio"),
        ("HA-25 B 20", "m3", "72,50"),
        ("Bombeo", "m3", "9,00"),
    ]
    t = doc.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            t.rows[r].cells[c].text = filas[r][c]
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@requiere_markitdown
def test_docx_conserva_tabla_de_tarifas() -> None:
    docx = _docx_con_tabla()
    if docx is None:
        pytest.skip("python-docx no instalado")
    md = a_markdown(docx, "contrato_0464.docx")
    assert md is not None
    assert "CONTRATO 0464" in md
    # La tabla debe venir como tabla Markdown (pipes + fila separadora).
    assert "| HA-25 B 20 |" in md
    assert "| --- |" in md or "---" in md


@requiere_markitdown
def test_combinar_varias_fuentes_respeta_orden_y_titulos() -> None:
    docx = _docx_con_tabla()
    if docx is None:
        pytest.skip("python-docx no instalado")
    fuentes = [
        FuenteDocumento(titulo="DOC 1 (antiguo)", filename="c1.docx", contenido=docx),
        FuenteDocumento(titulo="DOC 2 (moderno)", filename="c2.docx", contenido=docx),
    ]
    combinado = combinar_a_markdown(fuentes)
    assert combinado is not None
    assert combinado.index("DOC 1 (antiguo)") < combinado.index("DOC 2 (moderno)")
    assert "## DOC 1 (antiguo)" in combinado
    assert "\n---\n" in combinado  # separador entre secciones


def test_degradacion_usa_fallback_si_no_convierte() -> None:
    # Bytes basura con extensión .doc (binario antiguo que markitdown no
    # convierte): debe caer al texto de fallback, no lanzar.
    md = a_markdown(
        b"\x00\x01basura-no-docx",
        "viejo.doc",
        texto_fallback="texto plano del contrato",
    )
    assert md == "texto plano del contrato"


def test_documento_vacio_devuelve_fallback_o_none() -> None:
    assert a_markdown(b"", "x.pdf") is None
    assert a_markdown(b"", "x.pdf", texto_fallback="t") == "t"


def test_combinar_sin_secciones_devuelve_none() -> None:
    fuentes = [FuenteDocumento(titulo="t", filename="x.pdf", contenido=b"")]
    assert combinar_a_markdown(fuentes) is None
