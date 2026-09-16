# tests/test_f045_r15_r16_r17_escritura.py
"""F-045 · R15, R16 y R17: escribir los libros sin romper lo que ya había.

Los libros los rellena una persona y NO se versionan: si el importador los
estropea, no hay `git checkout` que valga. De ahí la copia previa (R16).

Y de ahí la regla que más código cuesta y más desastre evita (R17): el
importador actualiza las filas de SUS casos, pero **no degrada a `?` una celda
que ya tenía valor afirmado** ni borra las filas que él no genera. El banco ya
tenía 7 casos escritos a mano con mucho más detalle del que cabe en la tabla
plana —el número real del albarán, el `match_method` semántico, el volumen en
m3—, y los 7 están también en el Excel: un volcado a pelo los habría barrido
en la primera pasada.
"""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from evals.revision import escritura


def libro_de(tmp_path, filas, nombre="IA1_extraccion.xlsx"):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Residuos"
    hoja.append(["TABLA 1 — CABECERAS (una fila por albarán)"])
    hoja.append(["caso_id", "proveedor_nombre", "numero_albaran", "comentario"])
    for fila in filas:
        hoja.append(fila)
    hoja.append([])
    hoja.append(["TABLA 2 — LÍNEAS (una fila por línea)"])
    hoja.append(["caso_id", "num_linea", "cantidad", "comentario"])
    ruta = tmp_path / nombre
    libro.save(ruta)
    libro.close()
    return ruta


TABLAS = (
    escritura.DefTabla("TABLA 1", "cabeceras"),
    escritura.DefTabla("TABLA 2", "lineas"),
)


# --- R16: copia antes de escribir ------------------------------------------


def test_f045_r16_se_copia_el_libro_antes_de_tocarlo(tmp_path):
    ruta = libro_de(tmp_path, [["RES-001", "SALMEDINA", "SS-0000168", ""]])
    copia = escritura.copia_de_seguridad(ruta, dt.datetime(2026, 9, 15, 18, 30))
    assert copia.name == "IA1_extraccion.xlsx.20260915-1830.xlsx"
    assert copia.parent.name == "copias"
    assert copia.read_bytes() == ruta.read_bytes()


def test_f045_r16_copiar_un_libro_que_no_existe_lo_dice(tmp_path):
    with pytest.raises(escritura.ErrorEscritura):
        escritura.copia_de_seguridad(tmp_path / "no-esta.xlsx")


# --- R15: las filas de título no se mueven ---------------------------------


def test_f045_r15_las_filas_de_titulo_siguen_donde_estaban(tmp_path):
    ruta = libro_de(tmp_path, [])
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "X"}],
         "lineas": [{"caso_id": "RES-008", "num_linea": 1, "cantidad": 3}]},
        {"RES-008"},
    )
    hoja = openpyxl.load_workbook(ruta)["Residuos"]
    titulos = [c[0].value for c in hoja.iter_rows() if isinstance(c[0].value, str)
               and str(c[0].value).startswith("TABLA")]
    assert titulos == ["TABLA 1 — CABECERAS (una fila por albarán)",
                       "TABLA 2 — LÍNEAS (una fila por línea)"]


def test_f045_r15_la_fila_nueva_se_escribe_bajo_sus_encabezados(tmp_path):
    ruta = libro_de(tmp_path, [])
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "SALMEDINA",
                        "numero_albaran": "?", "comentario": None}],
         "lineas": []},
        {"RES-008"},
    )
    hoja = openpyxl.load_workbook(ruta)["Residuos"]
    assert [c.value for c in hoja[3]] == ["RES-008", "SALMEDINA", "?", None]


# --- R17: no se pisa lo que escribió el humano -----------------------------


def test_f045_r17_las_filas_de_otro_caso_no_se_tocan(tmp_path):
    ruta = libro_de(tmp_path, [["RES-001", "SALMEDINA", "SS-0000168", "a mano"]])
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "OTRO"}], "lineas": []},
        {"RES-008"},
    )
    hoja = openpyxl.load_workbook(ruta)["Residuos"]
    assert [c.value for c in hoja[3]] == ["RES-001", "SALMEDINA", "SS-0000168", "a mano"]


def test_f045_r17_el_interrogante_no_borra_un_valor_ya_afirmado(tmp_path):
    """El importador no sabe el número impreso; el humano sí lo escribió."""
    ruta = libro_de(tmp_path, [["RES-001", "SALMEDINA", "SS-0000168", "a mano"]])
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-001", "proveedor_nombre": "SALMEDINA, S.L.",
                        "numero_albaran": "?", "comentario": None}],
         "lineas": []},
        {"RES-001"},
    )
    hoja = openpyxl.load_workbook(ruta)["Residuos"]
    # El dato nuevo sí entra; el `?` deja en pie lo que ya estaba afirmado.
    assert [c.value for c in hoja[3]] == ["RES-001", "SALMEDINA, S.L.", "SS-0000168", None]


def test_f045_r17_una_fila_que_el_importador_no_genera_se_conserva(tmp_path):
    ruta = libro_de(tmp_path, [])
    libro = openpyxl.load_workbook(ruta)
    hoja = libro["Residuos"]
    hoja.insert_rows(7)
    for col, valor in enumerate(["RES-001", 1, 6, "volumen a mano"], start=1):
        hoja.cell(row=7, column=col, value=valor)
    libro.save(ruta)
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [], "lineas": [{"caso_id": "RES-001", "num_linea": 2, "cantidad": 9}]},
        {"RES-001"},
    )
    hoja = openpyxl.load_workbook(ruta)["Residuos"]
    filas = [[c.value for c in f] for f in hoja.iter_rows(min_row=6) if f[0].value]
    assert ["RES-001", 1, 6, "volumen a mano"] in filas
    assert ["RES-001", 2, 9, None] in filas


# --- R18: dos pasadas dejan el mismo libro ---------------------------------


def test_f045_r18_dos_pasadas_dejan_exactamente_el_mismo_contenido(tmp_path):
    ruta = libro_de(tmp_path, [["RES-001", "SALMEDINA", "SS-0000168", ""]])
    datos = {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "OTRO",
                            "numero_albaran": "?", "comentario": None}],
             "lineas": [{"caso_id": "RES-008", "num_linea": 1, "cantidad": 3}]}

    def contenido():
        hoja = openpyxl.load_workbook(ruta)["Residuos"]
        return [[c.value for c in fila] for fila in hoja.iter_rows()]

    escritura.escribir_pestana(ruta, "Residuos", TABLAS, datos, {"RES-008"})
    primera = contenido()
    escritura.escribir_pestana(ruta, "Residuos", TABLAS, datos, {"RES-008"})
    assert contenido() == primera


def test_f045_r15_una_pestana_sin_la_tabla_declarada_lo_dice(tmp_path):
    ruta = libro_de(tmp_path, [])
    with pytest.raises(escritura.ErrorEscritura) as fallo:
        escritura.escribir_pestana(
            ruta, "Residuos", (escritura.DefTabla("TABLA 9", "inventada"),),
            {"inventada": []}, set(),
        )
    assert "TABLA 9" in str(fallo.value)
