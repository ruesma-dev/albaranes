# evals/revision/escritura.py
"""Escribe los libros de ground truth. Infraestructura: aquí vive openpyxl.

Tres cosas que este módulo hace a propósito:

1. **Copia antes de escribir** (R16). Los libros los rellena una persona y no
   se versionan: si el importador los estropea, no hay `git checkout` que
   valga. La copia va a `ground_truth/copias/` con la fecha en el nombre.
2. **Las tablas se localizan por su fila de título**, no por coordenadas: el
   humano añade y quita filas y las posiciones se mueven entre una
   importación y la siguiente. Mismo criterio que `evals/conversor.py`, que es
   quien luego las lee; si las dos declaraciones divergen, el importador
   escribe donde el conversor no mira.
3. **Fusión conservadora, no volcado** (R8, R17). El importador actualiza las
   filas de SUS casos, pero nunca degrada a `?` una celda que ya tenía valor
   afirmado y deja intactas las filas que él no genera. El banco ya tenía 7
   casos escritos a mano con mucho más detalle del que cabe en la tabla plana
   —el número real del albarán, el `match_method` semántico, el volumen en
   m3—, y los 7 están también en el Excel: un volcado a pelo los habría
   barrido en la primera pasada.
"""

from __future__ import annotations

import datetime as dt
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from evals.conversor import clave_de_encabezado, sin_acentos
from evals.revision.vocabulario import cargar as cargar_vocabulario
from evals.revision.vocabulario import normalizar

#: Dónde se deja la copia de seguridad previa (R16).
NOMBRE_COPIAS = "copias"

#: Pestaña de la que se copia la estructura de una tipología nueva.
PLANTILLA = "Generico-Suministros"

#: Campos que identifican una fila dentro de su tabla. Es lo que permite
#: ACTUALIZAR la fila de un caso en vez de duplicarla, y lo que distingue una
#: fila del importador de otra que escribió el humano a mano.
CLAVES_DE_TABLA: dict[str, tuple[str, ...]] = {
    "cabeceras": ("caso_id",),
    "lineas": ("caso_id", "num_linea"),
    "contexto": ("caso_id", "num_linea", "campo_contexto"),
    "lineas_valoradas": ("caso_id", "num_linea"),
    "sinteticas_esperadas": ("caso_id", "num_linea_base", "descripcion_esperada"),
    "sinteticas_prohibidas": ("caso_id", "concepto_vetado"),
    "conciliacion": ("caso_id", "num_linea"),
    "caso": ("caso_id",),
    "lineas_albaran": ("caso_id", "num_linea"),
    "contrato_lineas": ("caso_id", "codigo_producto"),
    "condiciones": ("caso_id", "campo"),
    "datos_generales": ("caso_id",),
    "lineas_anadidas": ("caso_id", "num_linea_base", "concepto"),
}


#: Tablas cuya identidad incluye un texto libre que el humano y el Excel
#: escriben con distinto detalle. Ahí la descripción casa por PREFIJO: el libro
#: dice `INCREMENTO LER 170604` y la revisión trae el concepto entero, y son la
#: misma sintética.
CLAVE_POR_PREFIJO: dict[str, str] = {
    "sinteticas_esperadas": "descripcion_esperada",
    "lineas_anadidas": "concepto",
}


class ErrorEscritura(RuntimeError):
    """El libro no tiene la forma que el contrato de datos declara."""


@dataclass(frozen=True)
class DefTabla:
    """Una tabla dentro de una pestaña: por qué título empieza y cómo se llama."""

    titulo: str
    clave: str


@dataclass
class Bloque:
    """Dónde empieza y acaba una tabla ya localizada en la hoja."""

    definicion: DefTabla
    fila_encabezados: int = 0
    encabezados: list[str] | None = None
    primera_fila: int = 0
    ultima_fila: int = 0


# --- Copia previa y pestañas nuevas ----------------------------------------


def copia_de_seguridad(ruta: Path | str, momento: dt.datetime | None = None) -> Path:
    """`IA1_extraccion.xlsx` → `copias/IA1_extraccion.xlsx.20260915-1830.xlsx`."""
    origen = Path(ruta)
    if not origen.is_file():
        raise ErrorEscritura(f"no existe el libro '{origen}'.")
    sello = (momento or dt.datetime.now()).strftime("%Y%m%d-%H%M")
    destino = origen.parent / NOMBRE_COPIAS / f"{origen.name}.{sello}.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(origen, destino)
    return destino


def asegurar_pestanas(
    libro, pestanas: tuple[str, ...], plantilla: str = PLANTILLA
) -> list[str]:
    """Crea las pestañas que falten copiando la estructura de la plantilla.

    `Grava` y `Ferreteria` no existían en los libros: se crean copiando
    `Generico-Suministros`, que trae las mismas tablas y encabezados y ningún
    caso. Sin plantilla NO se inventa nada: una pestaña de estructura
    improvisada rompería el conversor más tarde y más lejos.
    """
    creadas: list[str] = []
    for pestana in pestanas:
        if pestana in libro.sheetnames:
            continue
        if plantilla not in libro.sheetnames:
            raise ErrorEscritura(
                f"falta la pestaña '{pestana}' y no está la plantilla "
                f"'{plantilla}' de la que copiar su estructura."
            )
        copia = libro.copy_worksheet(libro[plantilla])
        copia.title = pestana
        creadas.append(pestana)
    return creadas


# --- Localización de las tablas dentro de una pestaña (R15) ----------------


def _titulo_de(valor: object, titulos: tuple[str, ...]) -> str | None:
    if not isinstance(valor, str):
        return None
    normalizado = sin_acentos(valor).strip().upper()
    for titulo in titulos:
        if normalizado.startswith(sin_acentos(titulo).upper()):
            return titulo
    return None


def localizar_bloques(hoja, definiciones: tuple[DefTabla, ...]) -> list[Bloque]:
    """Encuentra las tablas por su fila de TÍTULO, nunca por coordenadas."""
    titulos = tuple(definicion.titulo for definicion in definiciones)
    por_titulo = {definicion.titulo: definicion for definicion in definiciones}
    bloques: list[Bloque] = []
    actual: Bloque | None = None

    for fila in hoja.iter_rows():
        valores = [celda.value for celda in fila]
        titulo = _titulo_de(valores[0] if valores else None, titulos)
        if titulo is not None:
            if actual is not None:
                actual.ultima_fila = fila[0].row - 1
            actual = Bloque(por_titulo[titulo])
            bloques.append(actual)
            continue
        if actual is None or actual.encabezados is not None:
            continue
        if all(valor is None for valor in valores):
            continue
        actual.encabezados = [
            clave_de_encabezado(valor) for valor in valores if valor is not None
        ]
        actual.fila_encabezados = fila[0].row
        actual.primera_fila = fila[0].row + 1
        actual.ultima_fila = max(hoja.max_row, fila[0].row)

    ausentes = [
        definicion.titulo
        for definicion in definiciones
        if definicion.titulo not in {bloque.definicion.titulo for bloque in bloques}
    ]
    if ausentes:
        raise ErrorEscritura(
            f"la pestaña '{hoja.title}' no tiene la fila de título de "
            f"{', '.join(ausentes)}. El contrato de datos la declara."
        )
    sin_encabezados = [b.definicion.titulo for b in bloques if b.encabezados is None]
    if sin_encabezados:
        raise ErrorEscritura(
            f"la pestaña '{hoja.title}': {', '.join(sin_encabezados)} no tiene "
            f"fila de encabezados debajo de su título."
        )
    return bloques


# --- Fusión conservadora (R8, R17) -----------------------------------------


def _clave(registro: dict, campos: tuple[str, ...]) -> tuple:
    return tuple(str(registro.get(campo, "")).strip() for campo in campos)


def _tiene_valor(valor: object) -> bool:
    if valor is None:
        return False
    texto = str(valor).strip()
    return bool(texto) and texto != "?"


def fundir(existente: dict, nuevo: dict) -> dict:
    """El importador actualiza, pero NO degrada un valor ya afirmado.

    Un `?` significa «no lo sé»; si el libro ya traía un valor escrito a mano,
    el que no sabe es el importador. Las columnas que no produce se quedan
    como están: la tabla plana no cubre todo lo que cabe en los libros.
    """
    fundido = dict(existente)
    for campo, valor in nuevo.items():
        if valor == "?" and _tiene_valor(existente.get(campo)):
            continue
        fundido[campo] = valor
    return fundido


def _casa_por_prefijo(
    existente: dict,
    nuevas: list[dict],
    campos_clave: tuple[str, ...],
    campo: str,
    existentes: list[dict],
) -> dict | None:
    """La misma línea descrita más corto o más largo (F-045, 2026-09-16).

    El libro escrito a mano dice `INCREMENTO LER 170604` y el Excel trae el
    concepto entero, `INCREMENTO LER 170604 MATERIALES DE AISLAMIENTO-E`. Son
    la MISMA sintética: escribir las dos haría que el banco esperase dos donde
    el sistema emite una, y ese rojo no existe.

    La pareja tiene que ser única **por los dos lados**: si una fila nueva
    podría ser la continuación de dos existentes, o al revés, NO se elige.
    Adivinar cuál es sería peor que dejar la fila aparte, donde se ve.
    """
    resto = tuple(c for c in campos_clave if c != campo)
    propio = normalizar_concepto(existente.get(campo))
    if not propio:
        return None

    def emparejan(fila_a: dict, fila_b: dict) -> bool:
        return _clave(fila_a, resto) == _clave(fila_b, resto) and _prefijo_comun(
            normalizar_concepto(fila_a.get(campo)), normalizar_concepto(fila_b.get(campo))
        )

    candidatos = [nueva for nueva in nuevas if emparejan(existente, nueva)]
    if len(candidatos) != 1:
        return None
    rivales = [otra for otra in existentes if emparejan(candidatos[0], otra)]
    return candidatos[0] if len(rivales) == 1 else None


_DIGITOS_SUELTOS = re.compile(r"(?<=\d) (?=\d)")


def normalizar_concepto(texto: object | None, sinonimos: dict[str, str] | None = None) -> str:
    """El mismo concepto escrito de dos maneras tiene que salir igual.

    Junta los dígitos separados —`17 08 02` es el LER `170802`— y aplica los
    sinónimos declarados en `vocabulario.json`, que hoy son las erratas con las
    que el humano escribe lo mismo en el Excel y en los libros. **Solo decide
    si dos filas son la misma línea**; el texto que se escribe no se toca.
    """
    palabras = _DIGITOS_SUELTOS.sub("", normalizar(texto)).split()
    tabla = _sinonimos() if sinonimos is None else sinonimos
    return " ".join(tabla.get(palabra, palabra) for palabra in palabras)


def _sinonimos() -> dict[str, str]:
    global _CACHE_SINONIMOS
    if _CACHE_SINONIMOS is None:
        _CACHE_SINONIMOS = {
            normalizar(k): normalizar(v)
            for k, v in cargar_vocabulario().sinonimos_concepto.items()
        }
    return _CACHE_SINONIMOS


_CACHE_SINONIMOS: dict[str, str] | None = None


def _prefijo_comun(uno: str, otro: str) -> bool:
    return bool(uno) and bool(otro) and (uno.startswith(otro) or otro.startswith(uno))


def _solo_del_importador(existente: dict, interrogantes: set[str]) -> bool:
    """La fila NO lleva nada que el importador no pudiera haber escrito.

    Es la condición para poder retirarla. Hace falta porque la fusión mueve la
    clave: cuando el importador actualiza una fila del humano, esa fila pasa a
    tener la clave del importador y, a partir de ahí, la huella la daría por
    suya. Retirarla se llevaría por delante lo que el humano escribió —pasó el
    2026-09-16 con el `modifier_source` de tres casos RES—.

    Así que una fila solo se retira si TODAS sus columnas están vacías o en
    `?` allí donde el importador escribe `?`. Si alguna lleva un valor que él
    nunca pone, la fila no es solo suya y se queda.

    `interrogantes` son las columnas que el importador deja en `?` **en toda la
    importación**, no solo en las filas de esta pestaña: si se dedujeran de las
    filas de turno, una pestaña donde esta vez no escribe nada no podría
    retirar nada, que es justo cuando hay que hacerlo —el humano ha quitado esa
    línea del Excel—. Sin columnas declaradas no se retira nada: ante la duda,
    no se toca lo del humano.
    """
    if not interrogantes:
        return False
    return not any(_tiene_valor(existente.get(campo)) for campo in interrogantes)


def fundir_filas(
    existentes: list[dict],
    nuevas: list[dict],
    campos_clave: tuple[str, ...],
    campo_prefijo: str | None = None,
    mias: set[tuple[str, ...]] | None = None,
    interrogantes: set[str] | None = None,
) -> list[dict]:
    """Actualiza en su sitio, conserva lo ajeno y añade al final lo nuevo.

    `mias` son las claves que el importador escribió la vez anterior (ver
    `huella.py`). Una fila que está ahí y que ya no se produce **se retira**:
    es suya y ha dejado de tener sentido. Lo que no está en `mias` se conserva
    siempre, porque lo escribió el humano (R17). Sin `mias`, no se retira nada.
    """
    por_clave = {_clave(nueva, campos_clave): nueva for nueva in nuevas}
    columnas_en_duda = interrogantes if interrogantes is not None else {
        campo for fila in nuevas for campo, valor in fila.items() if valor == "?"
    }
    resultado: list[dict] = []
    usadas: set[tuple] = set()
    for existente in existentes:
        clave = _clave(existente, campos_clave)
        nueva = por_clave.get(clave)
        if nueva is None and campo_prefijo:
            nueva = _casa_por_prefijo(
                existente, nuevas, campos_clave, campo_prefijo, existentes
            )
            if nueva is not None and _clave(nueva, campos_clave) in usadas:
                nueva = None
        if nueva is None:
            if (
                mias is not None
                and clave in mias
                and _solo_del_importador(existente, columnas_en_duda)
            ):
                continue  # la escribió el importador y ya no la produce
            resultado.append(existente)
            continue
        resultado.append(fundir(existente, nueva))
        usadas.add(_clave(nueva, campos_clave))
    for nueva in nuevas:
        if _clave(nueva, campos_clave) not in usadas:
            resultado.append(nueva)
    return resultado


# --- Escritura de una pestaña ----------------------------------------------


def _vacia(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _filas_existentes(hoja, bloque: Bloque) -> list[dict]:
    ancho = len(bloque.encabezados or [])
    filas: list[dict] = []
    for fila in hoja.iter_rows(min_row=bloque.primera_fila, max_row=bloque.ultima_fila):
        valores = [celda.value for celda in fila[:ancho]]
        if all(_vacia(valor) for valor in valores):
            continue
        filas.append(dict(zip(bloque.encabezados or [], valores)))
    return filas


def escribir_pestana(
    ruta: Path | str,
    pestana: str,
    definiciones: tuple[DefTabla, ...],
    filas_por_tabla: dict[str, list[dict]],
    caso_ids: set[str],
    ejecutar: bool = True,
    huella_previa: dict[str, set[tuple[str, ...]]] | None = None,
    interrogantes: dict[str, set[str]] | None = None,
) -> bool:
    """Vuelca las tablas de una pestaña fundiéndolas con lo que ya había.

    `caso_ids` son los casos de ESTA importación: solo sus filas se tocan.
    Devuelve si la pestaña cambia (o cambiaría, con `ejecutar=False`).

    **Si nada cambia, el libro NO se guarda.** No es una optimización: cada
    guardado reescribe el `.zip` del `.xlsx` y le cambia el sha256, que es
    justo la huella con la que el conversor detecta deriva entre el libro y
    sus fixtures. Guardar por guardar haría que dos importaciones seguidas
    dejasen 264 fixtures «modificados» sin que hubiera cambiado ni un dato
    (R18).
    """
    import openpyxl

    camino = Path(ruta)
    libro = openpyxl.load_workbook(camino)
    try:
        if pestana not in libro.sheetnames:
            raise ErrorEscritura(f"falta la pestaña '{pestana}' en '{camino.name}'.")
        hoja = libro[pestana]
        cambia = False
        # De abajo arriba: insertar o borrar filas mueve todo lo que hay
        # debajo, y así los índices de las tablas de arriba siguen valiendo.
        for bloque in sorted(
            localizar_bloques(hoja, definiciones),
            key=lambda b: b.fila_encabezados,
            reverse=True,
        ):
            nuevas = [
                fila
                for fila in filas_por_tabla.get(bloque.definicion.clave, [])
                if not caso_ids or fila.get("caso_id") in caso_ids
            ]
            campos = CLAVES_DE_TABLA.get(bloque.definicion.clave, ("caso_id",))
            existentes = _filas_existentes(hoja, bloque)
            fundidas = fundir_filas(
                existentes, nuevas, campos,
                campo_prefijo=CLAVE_POR_PREFIJO.get(bloque.definicion.clave),
                mias=huella_previa.get(bloque.definicion.clave) if huella_previa else None,
                interrogantes=(interrogantes or {}).get(bloque.definicion.clave),
            )
            cambia = cambia or _difieren(existentes, fundidas, bloque)
            if ejecutar:
                _reescribir(hoja, bloque, fundidas)
        if ejecutar and cambia:
            libro.save(camino)
        return cambia
    finally:
        libro.close()


def _difieren(existentes: list[dict], fundidas: list[dict], bloque: Bloque) -> bool:
    """Compara solo las columnas que el libro tiene: lo demás no se escribe."""
    columnas = bloque.encabezados or []

    def recorte(filas: list[dict]) -> list[list]:
        return [[_a_celda(fila.get(columna)) for columna in columnas] for fila in filas]

    return recorte(existentes) != recorte(fundidas)


def _ultima_con_datos(hoja, bloque: Bloque) -> int:
    """La última fila del bloque que tiene algo escrito.

    Las filas en blanco del final NO se tocan: son el aire que el humano dejó
    entre una tabla y la siguiente, y comérselas subiría las filas de título
    de todas las tablas de abajo en cada importación (R15).
    """
    ancho = len(bloque.encabezados or [])
    ultima = bloque.primera_fila - 1
    for fila in hoja.iter_rows(min_row=bloque.primera_fila, max_row=bloque.ultima_fila):
        if not all(_vacia(celda.value) for celda in fila[:ancho]):
            ultima = fila[0].row
    return ultima


def _reescribir(hoja, bloque: Bloque, filas: list[dict]) -> None:
    """Sustituye el bloque de datos por las filas ya fundidas."""
    sobrantes = _ultima_con_datos(hoja, bloque) - bloque.primera_fila + 1
    if sobrantes > 0:
        hoja.delete_rows(bloque.primera_fila, sobrantes)
    if not filas:
        return
    hoja.insert_rows(bloque.primera_fila, len(filas))
    for desplazamiento, registro in enumerate(filas):
        for columna, campo in enumerate(bloque.encabezados or [], start=1):
            hoja.cell(
                row=bloque.primera_fila + desplazamiento,
                column=columna,
                value=_a_celda(registro.get(campo)),
            )


def _a_celda(valor: object) -> object | None:
    """Lo que se escribe en la celda. Las listas viajan como `a;b` (contrato)."""
    if isinstance(valor, (list, tuple)):
        return ";".join(str(trozo) for trozo in valor) or None
    return valor
