# tests/conftest.py
"""Utilidades compartidas por los tests de F-011.

Los tests NUNCA tocan los libros reales de `evals/ground_truth/`: construyen
libros sintéticos con la MISMA estructura (títulos de tabla y encabezados
literales, copiados de los libros que rellena el humano) en un directorio
temporal. Así el test comprueba el contrato de datos de verdad, y sigue
funcionando en una máquina donde esos libros no existen.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

#: Las siete tipologías que tienen pestaña propia en los libros por tipología.
TIPOLOGIAS: tuple[str, ...] = (
    "Generico-Suministros",
    "Hormigon",
    "Mortero",
    "Residuos",
    "Bombeo",
    "Combustible",
    "Alquiler",
)

_ENC_IA1_CABECERAS = [
    "caso_id",
    "fichero_albaran",
    "proveedor_nombre",
    "proveedor_cif",
    "fecha (AAAA-MM-DD)",
    "numero_albaran",
    "obra_codigo",
    "obra_nombre",
    "forma_pago",
    "comentario",
]
_ENC_IA1_LINEAS = [
    "caso_id",
    "num_linea",
    "descripcion_esperada",
    "cantidad",
    "unidad",
    "precio_unitario",
    "descuentos (a;b)",
    "importe",
    "codigo_imputacion",
    "comentario",
]
_ENC_IA2 = [
    "caso_id",
    "num_linea",
    "campo_contexto",
    "valor_esperado",
    "comentario",
]
_ENC_IA3_T1 = [
    "caso_id",
    "num_linea",
    "match_method",
    "codigo_producto_contrato",
    "codigo_partida_final",
    "precio_unitario_final",
    "precio_source",
    "importe_calculado",
    "review_required (SI/NO)",
    "comentario",
]
_ENC_IA3_T2 = [
    "caso_id",
    "num_linea_base",
    "modifier_source",
    "rol_linea",
    "descripcion_esperada",
    "cantidad",
    "precio_unitario",
    "codigo_partida (heredada)",
    "comentario",
]
_ENC_IA3_T3 = [
    "caso_id",
    "modifier_source / concepto vetado",
    "motivo del veto",
    "comentario",
]
_ENC_IA4 = [
    "caso_id",
    "num_linea",
    "concilia (SI/NO)",
    "linea_contrato_esperada",
    "precio_unitario_esperado",
    "regla/motivo (p.ej. año no casa)",
    "comentario",
]
_ENC_FINAL_T1 = [
    "caso_id",
    "fichero (en inputs/albaranes)",
    "obra (código)",
    "proveedor (razón social)",
    "CIF",
    "fecha",
    "nº albarán",
    "contrato elegido (código)",
    "total valorado esperado (€)",
    "¿requiere revisión humana? (SI/NO)",
    "si SI, ¿por qué?",
    "comentario",
]
_ENC_FINAL_T2 = [
    "caso_id",
    "nº línea",
    "descripción (la del albarán)",
    "cantidad final",
    "unidad final",
    "¿casa con contrato? (SI/NO)",
    "línea del contrato con la que casa (código o descripción)",
    "partida final (código, ALM o REVISIÓN)",
    "precio unitario final (€ o REVISIÓN)",
    "el precio sale de (contrato / albarán / PDF contrato)",
    "importe final (€)",
    "¿línea a revisión? (SI/NO) y por qué",
    "comentario",
]
_ENC_FINAL_T3 = [
    "caso_id",
    "sobre qué línea va (nº)",
    "concepto (ej. INCREMENTO AÑO 2025, CANON DE VERTEDERO)",
    "cantidad",
    "precio unitario (€ o REVISIÓN)",
    "partida (normalmente la de su línea base)",
    "importe (€)",
    "comentario",
]

_TIT_IA3_T1 = "TABLA 1 — LÍNEAS VALORADAS (una fila por línea del albarán)"
_TIT_IA3_T2 = "TABLA 2 — SINTÉTICAS ESPERADAS (una fila por sintética que DEBE emitirse)"
_TIT_IA3_T3 = (
    "TABLA 3 — SINTÉTICAS PROHIBIDAS (una fila por sintética que NO debe aparecer)"
)
_TIT_INPUTS_LINEAS = (
    "LÍNEAS DEL ALBARÁN (solo casos con origen=manual; una fila por línea)"
)
_TIT_FINAL_T1 = "TABLA 1 — DATOS GENERALES DEL ALBARÁN (una fila por albarán)"
_TIT_FINAL_T2 = (
    "TABLA 2 — LÍNEAS DEL ALBARÁN, COMO DEBEN QUEDAR (una fila por línea impresa)"
)
_TIT_FINAL_T3 = (
    "TABLA 3 — LÍNEAS AÑADIDAS QUE NO ESTÁN IMPRESAS (incrementos, canon, portes...)"
)

#: Estructura literal de los seis libros: por libro, sus pestañas y, en cada
#: pestaña, la lista de (título de tabla, encabezados).
ESTRUCTURA: dict[str, dict[str, list[tuple[str, list[str]]]]] = {
    "IA1_extraccion.xlsx": {
        tipologia: [
            ("TABLA 1 — CABECERAS (una fila por albarán)", _ENC_IA1_CABECERAS),
            ("TABLA 2 — LÍNEAS (una fila por línea del albarán)", _ENC_IA1_LINEAS),
        ]
        for tipologia in TIPOLOGIAS
    },
    "IA2_contexto.xlsx": {
        tipologia: [
            ("CONTEXTO ESPERADO (una fila por caso, línea y campo)", _ENC_IA2)
        ]
        for tipologia in TIPOLOGIAS
    },
    "IA3_valoracion.xlsx": {
        tipologia: [
            (_TIT_IA3_T1, _ENC_IA3_T1),
            (_TIT_IA3_T2, _ENC_IA3_T2),
            (_TIT_IA3_T3, _ENC_IA3_T3),
        ]
        for tipologia in TIPOLOGIAS
    },
    "IA4_conciliacion.xlsx": {
        tipologia: [
            ("CONCILIACIÓN ESPERADA (una fila por línea sin match)", _ENC_IA4)
        ]
        for tipologia in TIPOLOGIAS
    },
    "INPUTS.xlsx": {
        "CASOS": [
            (
                "CASOS (registro maestro, una fila por caso)",
                [
                    "caso_id",
                    "tipologia",
                    "ia_destino (IA3/IA4/ambas)",
                    "origen (manual | merge:<caso_id>)",
                    "contrato_codigo",
                    "descripcion_caso",
                ],
            )
        ],
        "LINEAS_ALBARAN": [
            (
                _TIT_INPUTS_LINEAS,
                [
                    "caso_id",
                    "num_linea",
                    "descripcion",
                    "cantidad",
                    "unidad",
                    "precio_unitario",
                    "descuentos (a;b)",
                    "importe",
                    "codigo_imputacion",
                    "observaciones_albaran",
                ],
            )
        ],
        "CONTRATO_LINEAS": [
            (
                "LÍNEAS DEL CONTRATO (una fila por recurso del contrato)",
                [
                    "caso_id",
                    "codigo_producto",
                    "descripcion_recurso",
                    "unidad",
                    "precio_unitario",
                    "codigo_partida",
                    "comentario",
                ],
            )
        ],
        "CONDICIONES": [
            (
                "CONDICIONES DEL CASO (formato campo/valor)",
                ["caso_id", "campo", "valor", "comentario"],
            )
        ],
    },
    "RESULTADO_FINAL.xlsx": {
        tipologia: [
            (_TIT_FINAL_T1, _ENC_FINAL_T1),
            (_TIT_FINAL_T2, _ENC_FINAL_T2),
            (_TIT_FINAL_T3, _ENC_FINAL_T3),
        ]
        for tipologia in TIPOLOGIAS
    },
}

#: Clave de `filas`: (fichero, pestaña, índice de tabla dentro de la pestaña).
Ubicacion = tuple[str, str, int]


def construir_ground_truth(
    directorio: Path,
    filas: dict[Ubicacion, list[list[object]]] | None = None,
    omitir_libros: tuple[str, ...] = (),
    omitir_pestanas: tuple[tuple[str, str], ...] = (),
) -> Path:
    """Escribe los seis libros sintéticos en `directorio` y lo devuelve.

    `filas` inyecta filas de datos en una tabla concreta; lo que no se inyecta
    queda como en los libros que hoy tiene el humano: solo títulos y cabeceras.
    `omitir_libros` y `omitir_pestanas` sirven para los tests de R6.
    """
    filas = filas or {}
    directorio.mkdir(parents=True, exist_ok=True)

    for fichero, pestanas in ESTRUCTURA.items():
        if fichero in omitir_libros:
            continue
        libro = openpyxl.Workbook()
        leeme = libro.active
        leeme.title = "LEEME"
        leeme["A1"] = f"GROUND TRUTH sintético de test — {fichero}"

        for pestana, tablas in pestanas.items():
            if (fichero, pestana) in omitir_pestanas:
                continue
            hoja = libro.create_sheet(pestana)
            fila_actual = 1
            for indice, (titulo, encabezados) in enumerate(tablas):
                hoja.cell(row=fila_actual, column=1, value=titulo)
                fila_actual += 1
                for columna, encabezado in enumerate(encabezados, start=1):
                    hoja.cell(row=fila_actual, column=columna, value=encabezado)
                fila_actual += 1
                for datos in filas.get((fichero, pestana, indice), []):
                    for columna, valor in enumerate(datos, start=1):
                        hoja.cell(row=fila_actual, column=columna, value=valor)
                    fila_actual += 1
                fila_actual += 2  # filas en blanco de separación, como en el real

        libro.save(directorio / fichero)

    return directorio


@pytest.fixture
def constructor_gt():
    """Da a los tests el constructor de libros sintéticos, sin importarlo.

    Se expone como fixture y no por `import`: los tests de la raíz y los de
    los servicios comparten el mismo `sys.path`, y un módulo de utilidades
    importado por su nombre es una colisión esperando a pasar.
    """
    return construir_ground_truth


@pytest.fixture
def ground_truth_vacio(tmp_path: Path) -> Path:
    """Los seis libros con estructura y sin ningún caso (el estado de hoy)."""
    return construir_ground_truth(tmp_path / "ground_truth")
