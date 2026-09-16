# tests/test_f045_r17_r18_fusion.py
"""F-045 · R17 y R18: la mecánica fina de la fusión y del «no guardar».

Segunda tanda de tests nacidos de la campaña de mutación. Aquí viven las
piezas pequeñas de `escritura.py` que deciden **qué filas se pisan y cuándo se
guarda el libro**, y son pequeñas justamente porque cada una responde a una
sola pregunta. Que ninguna estuviera comprobada por separado es lo que dejó
vivos 23 mutantes de ese fichero: el banco miraba el resultado de una pasada
feliz y no las decisiones que lo producen.
"""

from __future__ import annotations

import openpyxl
import pytest

from evals.revision import escritura

TABLAS = (
    escritura.DefTabla("TABLA 1", "cabeceras"),
    escritura.DefTabla("TABLA 2", "lineas"),
)


def libro_de(tmp_path, filas_t1=(), filas_t2=(), huecos=2):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Residuos"
    hoja.append(["TABLA 1 — CABECERAS"])
    hoja.append(["caso_id", "proveedor_nombre", "comentario"])
    for fila in filas_t1:
        hoja.append(list(fila))
    for _ in range(huecos):
        hoja.append([])
    hoja.append(["TABLA 2 — LÍNEAS"])
    hoja.append(["caso_id", "num_linea", "cantidad"])
    for fila in filas_t2:
        hoja.append(list(fila))
    ruta = tmp_path / "IA1_extraccion.xlsx"
    libro.save(ruta)
    libro.close()
    return ruta


def volcado(ruta, pestana="Residuos"):
    libro = openpyxl.load_workbook(ruta)
    filas = [[c.value for c in fila] for fila in libro[pestana].iter_rows()]
    libro.close()
    return filas


# --- `_tiene_valor`: qué cuenta como «ya afirmado» -----------------------


@pytest.mark.parametrize(
    "valor,esperado",
    [("SS-0000168", True), (0, True), ("?", False), ("", False), ("   ", False), (None, False)],
)
def test_f045_r17_solo_es_valor_afirmado_lo_que_dice_algo(valor, esperado):
    """`?` es «no lo sé»: un `?` en el libro no protege nada de la fusión."""
    assert escritura._tiene_valor(valor) is esperado


def test_f045_r17_el_interrogante_del_libro_si_lo_pisa_el_importador():
    existente = {"caso_id": "RES-001", "numero_albaran": "?"}
    nuevo = {"caso_id": "RES-001", "numero_albaran": "SS-0000168"}
    assert escritura.fundir(existente, nuevo)["numero_albaran"] == "SS-0000168"


def test_f045_r17_una_columna_que_el_importador_no_produce_se_queda():
    fundido = escritura.fundir({"caso_id": "RES-001", "forma_pago": "contado"},
                               {"caso_id": "RES-001"})
    assert fundido["forma_pago"] == "contado"


# --- La localización del bloque: dónde empieza y dónde acaba -------------


def test_f045_r15_el_bloque_empieza_justo_bajo_sus_encabezados(tmp_path):
    ruta = libro_de(tmp_path, filas_t1=[["RES-001", "SALMEDINA", None]])
    libro = openpyxl.load_workbook(ruta)
    cabeceras, lineas = escritura.localizar_bloques(libro["Residuos"], TABLAS)
    assert (cabeceras.fila_encabezados, cabeceras.primera_fila) == (2, 3)
    # El bloque de la primera tabla acaba en la fila ANTERIOR al título de la
    # segunda: comerse esa frontera pisaría la fila de título.
    assert cabeceras.ultima_fila == 5
    assert (lineas.fila_encabezados, lineas.primera_fila) == (7, 8)
    libro.close()


def test_f045_r15_las_filas_en_blanco_del_final_no_se_tocan(tmp_path):
    """Son el aire que el humano dejó entre tablas; comérselas subiría los

    títulos de todas las de abajo en cada importación."""
    ruta = libro_de(tmp_path, filas_t1=[["RES-001", "SALMEDINA", None]], huecos=3)
    antes = volcado(ruta)
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-001", "proveedor_nombre": "SALMEDINA, S.L."}],
         "lineas": []},
        {"RES-001"},
    )
    despues = volcado(ruta)
    assert len(despues) == len(antes)
    assert despues[6][0] == "TABLA 2 — LÍNEAS"


def test_f045_r15_una_tabla_sin_encabezados_bajo_su_titulo_lo_dice(tmp_path):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Residuos"
    hoja.append(["TABLA 1 — CABECERAS"])
    ruta = tmp_path / "sin.xlsx"
    libro.save(ruta)
    libro.close()
    libro = openpyxl.load_workbook(ruta)
    with pytest.raises(escritura.ErrorEscritura) as fallo:
        escritura.localizar_bloques(libro["Residuos"], (TABLAS[0],))
    assert "encabezados" in str(fallo.value)
    libro.close()


# --- Cuándo se guarda el libro y cuándo no (R18) -------------------------


def test_f045_r18_una_pestana_que_no_cambia_dice_que_no_cambia(tmp_path):
    ruta = libro_de(tmp_path, filas_t1=[["RES-001", "SALMEDINA", None]])
    filas = {"cabeceras": [{"caso_id": "RES-001", "proveedor_nombre": "SALMEDINA, S.L.",
                            "comentario": None}], "lineas": []}
    # La primera pasada corrige la razón social: eso sí es un cambio.
    assert escritura.escribir_pestana(ruta, "Residuos", TABLAS, filas, {"RES-001"}) is True
    # La segunda escribe lo mismo que ya hay: no hay nada que guardar.
    assert escritura.escribir_pestana(ruta, "Residuos", TABLAS, filas, {"RES-001"}) is False


def test_f045_r18_sin_cambios_el_fichero_no_se_reescribe(tmp_path):
    """Guardar un libro idéntico le mueve el sha256 y deja 264 fixtures
    «modificados» sin que haya cambiado un dato."""
    ruta = libro_de(tmp_path, filas_t1=[["RES-001", "SALMEDINA", None]])
    filas = {"cabeceras": [{"caso_id": "RES-001", "proveedor_nombre": "SALMEDINA",
                            "comentario": None}], "lineas": []}
    escritura.escribir_pestana(ruta, "Residuos", TABLAS, filas, {"RES-001"})
    antes = ruta.read_bytes()
    escritura.escribir_pestana(ruta, "Residuos", TABLAS, filas, {"RES-001"})
    assert ruta.read_bytes() == antes


def test_f045_r18_en_modo_comprobacion_no_se_escribe_nada(tmp_path):
    ruta = libro_de(tmp_path)
    antes = ruta.read_bytes()
    cambia = escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "X"}], "lineas": []},
        {"RES-008"}, ejecutar=False,
    )
    assert cambia is True
    assert ruta.read_bytes() == antes


# --- `caso_ids` acota qué filas entran ----------------------------------


def test_f045_r17_una_fila_de_un_caso_ajeno_a_la_importacion_no_entra(tmp_path):
    ruta = libro_de(tmp_path)
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "SI"},
                       {"caso_id": "RES-099", "proveedor_nombre": "NO"}],
         "lineas": []},
        {"RES-008"},
    )
    escritos = [f[0] for f in volcado(ruta) if f[0] and str(f[0]).startswith("RES-")]
    assert escritos == ["RES-008"]


def test_f045_r17_sin_caso_ids_entran_todas_las_filas(tmp_path):
    ruta = libro_de(tmp_path)
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008"}, {"caso_id": "RES-099"}], "lineas": []},
        set(),
    )
    escritos = [f[0] for f in volcado(ruta) if f[0] and str(f[0]).startswith("RES-")]
    assert escritos == ["RES-008", "RES-099"]


# --- Borrar exactamente las filas de datos, ni una más ------------------


def test_f045_r15_reescribir_borra_las_filas_de_datos_y_solo_esas(tmp_path):
    ruta = libro_de(
        tmp_path,
        filas_t1=[["RES-001", "UNO", None], ["RES-002", "DOS", None]],
        huecos=2,
    )
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-001", "proveedor_nombre": "UNO BIS",
                        "comentario": None}], "lineas": []},
        {"RES-001"},
    )
    filas = volcado(ruta)
    assert filas[0][0] == "TABLA 1 — CABECERAS"
    assert filas[1][0] == "caso_id"
    assert [f[:2] for f in filas[2:4]] == [["RES-001", "UNO BIS"], ["RES-002", "DOS"]]


def test_f045_r15_un_bloque_vacio_no_borra_la_fila_de_titulo_de_abajo(tmp_path):
    ruta = libro_de(tmp_path, huecos=1)
    antes = volcado(ruta)
    escritura.escribir_pestana(ruta, "Residuos", TABLAS,
                               {"cabeceras": [], "lineas": []}, set())
    assert volcado(ruta) == antes


# --- Listas a celda: el contrato dice `a;b` ------------------------------


def test_f045_r15_una_lista_viaja_como_texto_separado_por_punto_y_coma():
    assert escritura._a_celda(["10", "5"]) == "10;5"
    assert escritura._a_celda([]) is None
    assert escritura._a_celda("40") == "40"
    assert escritura._a_celda(None) is None


# --- R16: la copia es el estado PREVIO, o no es una copia de seguridad ---


def test_f045_r16_la_copia_es_el_libro_tal_como_estaba_antes_de_escribir(tmp_path):
    """El bloqueante que encontró el reviewer en la pasada 1.

    Si la comprobación previa —la que solo mira si el libro cambia— llegara a
    guardar, la copia se haría sobre el libro YA modificado y dejaría de ser lo
    que R16 promete: el estado al que volver. Y esa copia es la única red del
    trabajo manual del humano, porque los libros no se versionan (D5).
    """
    ruta = libro_de(tmp_path, filas_t1=[["RES-001", "SALMEDINA", "a mano"]])
    previo = ruta.read_bytes()

    cambia = escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "NUEVO"}],
         "lineas": []},
        {"RES-008"}, ejecutar=False,
    )
    assert cambia is True
    assert ruta.read_bytes() == previo, (
        "la pasada de comprobación ha escrito en el libro: la copia de "
        "seguridad que venga después ya no sería el estado previo (R16)"
    )

    copia = escritura.copia_de_seguridad(ruta)
    escritura.escribir_pestana(
        ruta, "Residuos", TABLAS,
        {"cabeceras": [{"caso_id": "RES-008", "proveedor_nombre": "NUEVO"}],
         "lineas": []},
        {"RES-008"},
    )
    assert copia.read_bytes() == previo, (
        f"la copia de {ruta.name} NO es el libro previo"
    )
    assert ruta.read_bytes() != previo  # y la escritura de verdad sí ocurrió
