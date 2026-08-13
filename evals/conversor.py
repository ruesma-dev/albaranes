# evals/conversor.py
"""Conversor de los libros de ground truth a fixtures JSON versionables.

Los seis libros Excel de `evals/ground_truth/` los rellena una persona y NO
entran en git (regla del arnés: la ofimática se queda fuera). Lo que se
versiona es su traducción: un JSON por caso, ordenado y estable, que es lo que
el runner compara contra el sistema.

Tres cosas que este módulo hace a propósito y conviene no «simplificar»:

1. **Todo en memoria antes de escribir.** El barrido de datos sensibles corre
   sobre el contenido completo y, si salta, no se ha creado ni un fichero. Unos
   fixtures a medias con un correo dentro ya estarían en el disco.
2. **Los sentinelas se escriben.** `?` y `REVISIÓN` viajan como valor explícito
   (ver `evals/modelos.py`), nunca como clave ausente: la ausencia no se
   distingue de un descuido del conversor.
3. **Salida determinista.** Claves ordenadas, UTF-8 sin escapes y ni un
   timestamp. La trazabilidad la da el sha256 del libro origen, que solo cambia
   cuando cambia el libro.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from fnmatch import fnmatchcase
from pathlib import Path

from evals.barrido import Hallazgo, barrer
from evals.modelos import ValorEsperado

#: Directorio de los libros que rellena el humano (no versionados).
RUTA_GROUND_TRUTH = Path(__file__).resolve().parent / "ground_truth"

#: Directorio de los fixtures generados (sí versionados).
RUTA_FIXTURES = Path(__file__).resolve().parent / "fixtures"

#: Fichero de índice que cada destino deja junto a sus casos. Da estructura
#: versionable aunque todavía no haya ni un caso, y deja a la vista el sha256
#: del libro del que salieron (detector de deriva Excel ↔ fixtures).
NOMBRE_INDICE = "_indice.json"

TIPOLOGIAS: tuple[str, ...] = (
    "Generico-Suministros",
    "Hormigon",
    "Mortero",
    "Residuos",
    "Bombeo",
    "Combustible",
    "Alquiler",
)

#: Campos cuyo valor de texto se interpreta como número si se puede. Fuera de
#: esta lista el texto se respeta tal cual: convertir «0012345» a 12345 en un
#: número de albarán sería estropear el ground truth.
CAMPOS_NUMERICOS: tuple[str, ...] = (
    "*cantidad*",
    "*precio*",
    "*importe*",
    "*total*",
    "*volumen*",
    "*peso*",
    "*num_linea*",
    "*litros*",
    "*horas*",
    "*m3*",
    "*contenedor*",
    "*confianza*",
)

#: Campos que el contrato de datos declara como lista separada por `;`.
CAMPOS_LISTA: tuple[str, ...] = ("descuentos",)

#: Encabezados escritos para personas cuya traducción automática a clave sería
#: ilegible o ambigua. El resto se derivan solos (ver `clave_de_encabezado`).
ALIAS_ENCABEZADOS: dict[str, str] = {
    "no_albaran": "numero_albaran",
    "no_linea": "num_linea",
    "sobre_que_linea_va": "num_linea_base",
    "linea_del_contrato_con_la_que_casa": "linea_contrato",
    "el_precio_sale_de": "precio_source",
    "linea_a_revision_y_por_que": "linea_a_revision",
    "requiere_revision_humana": "requiere_revision",
    "si_si_por_que": "motivo_revision",
    "modifier_source_concepto_vetado": "concepto_vetado",
    "motivo_del_veto": "motivo_veto",
    "regla_motivo": "motivo",
}

_PARENTESIS = re.compile(r"\([^)]*\)")
_NO_ALFANUMERICO = re.compile(r"[^0-9a-z]+")
_MILES_Y_COMA = re.compile(r"^-?\d{1,3}(?:\.\d{3})+(?:,\d+)?$")
_COMA_DECIMAL = re.compile(r"^-?\d+,\d+$")
_PUNTO_DECIMAL = re.compile(r"^-?\d+(?:\.\d+)?$")


class ErrorConversion(RuntimeError):
    """El contrato de datos no se cumple: no se convierte nada."""


class ErrorDatosSensibles(ErrorConversion):
    """El barrido de C3 bis encontró algo que no puede entrar en git."""

    def __init__(self, hallazgos: list[Hallazgo]) -> None:
        self.hallazgos = hallazgos
        detalle = "\n".join(f"  - {h.descripcion()}" for h in hallazgos)
        super().__init__(
            f"Barrido de datos sensibles: {len(hallazgos)} hallazgo(s). No se "
            f"ha escrito ningún fixture.\n{detalle}"
        )


# --- Declaración del contrato de datos --------------------------------------


@dataclass(frozen=True)
class DefTabla:
    """Una tabla dentro de una pestaña: por qué título empieza y cómo se llama."""

    titulo: str
    clave: str


@dataclass(frozen=True)
class DefLibro:
    """Un libro del ground truth y qué se espera encontrar dentro."""

    fichero: str
    fase: str
    destino: str
    tablas_por_pestana: dict[str, tuple[DefTabla, ...]]
    #: `True` cuando cada pestaña es una tipología (y por tanto da la tipología
    #: del caso); `False` cuando las pestañas son partes de un mismo caso.
    por_tipologia: bool = True


def _por_tipologia(tablas: tuple[DefTabla, ...]) -> dict[str, tuple[DefTabla, ...]]:
    return {tipologia: tablas for tipologia in TIPOLOGIAS}


#: Los seis libros del contrato de datos (ver `evals/README.md`).
LIBROS: tuple[DefLibro, ...] = (
    DefLibro(
        fichero="IA1_extraccion.xlsx",
        fase="IA1",
        destino="IA1",
        tablas_por_pestana=_por_tipologia(
            (DefTabla("TABLA 1", "cabeceras"), DefTabla("TABLA 2", "lineas"))
        ),
    ),
    DefLibro(
        fichero="IA2_contexto.xlsx",
        fase="IA2",
        destino="IA2",
        tablas_por_pestana=_por_tipologia(
            (DefTabla("CONTEXTO ESPERADO", "contexto"),)
        ),
    ),
    DefLibro(
        fichero="IA3_valoracion.xlsx",
        fase="IA3",
        destino="IA3",
        tablas_por_pestana=_por_tipologia(
            (
                DefTabla("TABLA 1", "lineas_valoradas"),
                DefTabla("TABLA 2", "sinteticas_esperadas"),
                DefTabla("TABLA 3", "sinteticas_prohibidas"),
            )
        ),
    ),
    DefLibro(
        fichero="IA4_conciliacion.xlsx",
        fase="IA4",
        destino="IA4",
        tablas_por_pestana=_por_tipologia(
            (DefTabla("CONCILIACIÓN ESPERADA", "conciliacion"),)
        ),
    ),
    DefLibro(
        fichero="INPUTS.xlsx",
        fase="INPUTS",
        destino="inputs",
        por_tipologia=False,
        tablas_por_pestana={
            "CASOS": (DefTabla("CASOS", "caso"),),
            "LINEAS_ALBARAN": (DefTabla("LÍNEAS DEL ALBARÁN", "lineas_albaran"),),
            "CONTRATO_LINEAS": (DefTabla("LÍNEAS DEL CONTRATO", "contrato_lineas"),),
            "CONDICIONES": (DefTabla("CONDICIONES DEL CASO", "condiciones"),),
        },
    ),
    DefLibro(
        fichero="RESULTADO_FINAL.xlsx",
        fase="FINAL",
        destino="final",
        tablas_por_pestana=_por_tipologia(
            (
                DefTabla("TABLA 1", "datos_generales"),
                DefTabla("TABLA 2", "lineas"),
                DefTabla("TABLA 3", "lineas_anadidas"),
            )
        ),
    ),
)


# --- Normalización de encabezados y celdas ----------------------------------


def sin_acentos(texto: str) -> str:
    """Quita tildes y signos compatibles (`nº` → `no`) sin tocar el resto."""
    descompuesto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in descompuesto if not unicodedata.combining(c))


def clave_de_encabezado(encabezado: str) -> str:
    """Traduce un encabezado escrito para personas a una clave estable.

    «precio unitario final (€ o REVISIÓN)» → `precio_unitario_final`. Lo que
    va entre paréntesis es ayuda para quien rellena el libro, no parte del
    nombre del campo.
    """
    texto = _PARENTESIS.sub(" ", str(encabezado))
    texto = sin_acentos(texto).lower()
    clave = _NO_ALFANUMERICO.sub("_", texto).strip("_")
    return ALIAS_ENCABEZADOS.get(clave, clave)


def _es_campo_numerico(campo: str) -> bool:
    return any(fnmatchcase(campo, patron) for patron in CAMPOS_NUMERICOS)


def _a_numero(texto: str) -> float | None:
    """Convierte «1.234,56», «72,50» o «72.5» a float; `None` si no es número."""
    limpio = texto.replace("€", "").replace(" ", "").strip()
    if _MILES_Y_COMA.match(limpio):
        limpio = limpio.replace(".", "").replace(",", ".")
    elif _COMA_DECIMAL.match(limpio):
        limpio = limpio.replace(",", ".")
    elif not _PUNTO_DECIMAL.match(limpio):
        return None
    try:
        return float(limpio)
    except ValueError:  # pragma: no cover - las regex ya lo garantizan
        return None


def normalizar_celda(valor: object, campo: str = "") -> ValorEsperado:
    """Aplica los convenios de celda del contrato de datos (R2)."""
    if isinstance(valor, dt.datetime):
        return ValorEsperado(valor.date().isoformat())
    if isinstance(valor, dt.date):
        return ValorEsperado(valor.isoformat())
    if valor is None:
        return ValorEsperado(None)
    if isinstance(valor, bool):
        return ValorEsperado(valor)
    if isinstance(valor, (int, float)):
        return ValorEsperado(valor)

    texto = str(valor).strip()
    if not texto:
        return ValorEsperado(None)
    if texto == "?":
        return ValorEsperado(None, comparar=False)
    if sin_acentos(texto).upper() == "REVISION":
        return ValorEsperado(None, espera_revision=True)
    if campo in CAMPOS_LISTA:
        return ValorEsperado([trozo.strip() for trozo in texto.split(";") if trozo.strip()])
    if _es_campo_numerico(campo):
        numero = _a_numero(texto)
        if numero is not None:
            return ValorEsperado(numero)
    return ValorEsperado(texto)


# --- Lectura de una pestaña -------------------------------------------------


@dataclass
class Tabla:
    """Una tabla ya leída: sus claves de columna y sus filas normalizadas."""

    clave: str
    encabezados: list[str] = field(default_factory=list)
    filas: list[dict[str, object | None]] = field(default_factory=list)
    caso_por_fila: list[str] = field(default_factory=list)


def _empieza_por_titulo(texto: object, titulos: tuple[str, ...]) -> str | None:
    if not isinstance(texto, str):
        return None
    normalizado = sin_acentos(texto).strip().upper()
    for titulo in titulos:
        if normalizado.startswith(sin_acentos(titulo).upper()):
            return titulo
    return None


def parsear_hoja(
    hoja,  # openpyxl Worksheet
    definiciones: tuple[DefTabla, ...],
    ubicacion: str,
    hallazgos: list[Hallazgo],
) -> dict[str, Tabla]:
    """Localiza en la pestaña las tablas declaradas y devuelve sus filas.

    Las tablas se buscan por su fila de título, no por posición: el humano
    añade y quita filas dentro de cada tabla y las coordenadas se mueven.
    """
    titulos = tuple(definicion.titulo for definicion in definiciones)
    por_titulo = {definicion.titulo: definicion for definicion in definiciones}

    tablas: dict[str, Tabla] = {}
    definicion_actual: DefTabla | None = None
    esperando_encabezados = False

    for fila in hoja.iter_rows():
        celdas = list(fila)
        primera = celdas[0].value if celdas else None
        titulo = _empieza_por_titulo(primera, titulos)

        if titulo is not None:
            definicion_actual = por_titulo[titulo]
            tablas[definicion_actual.clave] = Tabla(clave=definicion_actual.clave)
            esperando_encabezados = True
            continue

        if definicion_actual is None:
            continue

        tabla = tablas[definicion_actual.clave]
        if esperando_encabezados:
            if all(celda.value is None for celda in celdas):
                continue
            tabla.encabezados = [
                clave_de_encabezado(celda.value)
                for celda in celdas
                if celda.value is not None
            ]
            esperando_encabezados = False
            continue

        ancho = len(tabla.encabezados)
        valores = [celda.value for celda in celdas[:ancho]]
        if all(_esta_vacia(valor) for valor in valores):
            continue

        registro: dict[str, object | None] = {}
        for indice, campo in enumerate(tabla.encabezados):
            bruto = valores[indice] if indice < len(valores) else None
            if isinstance(bruto, str):
                hallazgos.extend(
                    barrer(
                        bruto,
                        ubicacion=f"{ubicacion}!{celdas[indice].coordinate}",
                    )
                )
            registro[campo] = normalizar_celda(bruto, campo).a_json()

        caso_id = registro.get("caso_id")
        if not isinstance(caso_id, str) or not caso_id.strip():
            raise ErrorConversion(
                f"{ubicacion}, fila {celdas[0].row}: hay datos pero la columna "
                f"'caso_id' está vacía. Cada fila pertenece a un caso; sin "
                f"caso_id no se sabe a cuál."
            )
        tabla.filas.append(registro)
        tabla.caso_por_fila.append(caso_id.strip())

    ausentes = [
        definicion.titulo
        for definicion in definiciones
        if definicion.clave not in tablas
    ]
    if ausentes:
        raise ErrorConversion(
            f"{ubicacion}: no se encuentra la fila de título de "
            f"{', '.join(ausentes)}. El contrato de datos declara "
            f"{len(definiciones)} tabla(s) en esta pestaña."
        )
    return tablas


def _esta_vacia(valor: object) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


# --- Conversión de un libro -------------------------------------------------


@dataclass
class CasoConvertido:
    """Un caso ya montado, a la espera de que el barrido dé permiso a escribir."""

    caso_id: str
    fase: str
    tipologia: str
    libro: str
    sha256_libro: str
    tablas: dict[str, list[dict[str, object | None]]]

    def a_json(self) -> dict[str, object]:
        return {
            "caso_id": self.caso_id,
            "fase": self.fase,
            "libro": self.libro,
            "sha256_libro": self.sha256_libro,
            "tablas": self.tablas,
            "tipologia": self.tipologia,
        }


@dataclass
class InformeConversion:
    """Qué se convirtió, cuántos casos por fase y qué encontró el barrido."""

    libros: list[str] = field(default_factory=list)
    casos_por_fase: dict[str, int] = field(default_factory=dict)
    ficheros: list[Path] = field(default_factory=list)
    hallazgos: list[Hallazgo] = field(default_factory=list)

    def resumen(self) -> str:
        detalle = ", ".join(
            f"{fase}: {total}" for fase, total in sorted(self.casos_por_fase.items())
        )
        return (
            f"{len(self.libros)} libro(s) convertidos, "
            f"{len(self.ficheros)} fichero(s) escritos ({detalle or 'sin casos'})"
        )


def sha256_de(ruta: Path) -> str:
    """Huella del libro origen: la trazabilidad que no depende del reloj."""
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def _convertir_libro(
    definicion: DefLibro, ruta: Path, hallazgos: list[Hallazgo]
) -> tuple[str, list[CasoConvertido]]:
    import openpyxl  # import perezoso: solo el conversor necesita openpyxl

    if not ruta.is_file():
        raise ErrorConversion(
            f"falta el libro '{definicion.fichero}' en {ruta.parent}. El "
            f"contrato de datos declara los {len(LIBROS)} libros de "
            f"evals/ground_truth/."
        )

    huella = sha256_de(ruta)
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=False)

    tablas_por_caso: dict[str, dict[str, list[dict[str, object | None]]]] = {}
    tipologia_de_caso: dict[str, str] = {}
    claves = [
        definicion_tabla.clave
        for tablas in definicion.tablas_por_pestana.values()
        for definicion_tabla in tablas
    ]

    for pestana, definiciones in definicion.tablas_por_pestana.items():
        if pestana not in libro.sheetnames:
            raise ErrorConversion(
                f"falta la pestaña '{pestana}' en el libro "
                f"'{definicion.fichero}'. El contrato de datos la declara."
            )
        ubicacion = f"{definicion.fichero}!{pestana}"
        tablas = parsear_hoja(libro[pestana], definiciones, ubicacion, hallazgos)

        for clave, tabla in tablas.items():
            for registro, caso_id in zip(tabla.filas, tabla.caso_por_fila):
                caso = tablas_por_caso.setdefault(
                    caso_id, {nombre: [] for nombre in claves}
                )
                caso[clave].append(registro)
                if definicion.por_tipologia:
                    tipologia_de_caso.setdefault(caso_id, pestana)
                elif clave == "caso" and isinstance(registro.get("tipologia"), str):
                    tipologia_de_caso[caso_id] = str(registro["tipologia"])

    libro.close()

    return huella, [
        CasoConvertido(
            caso_id=caso_id,
            fase=definicion.fase,
            tipologia=tipologia_de_caso.get(caso_id, ""),
            libro=definicion.fichero,
            sha256_libro=huella,
            tablas=tablas,
        )
        for caso_id, tablas in sorted(tablas_por_caso.items())
    ]


def _volcar(datos: object) -> str:
    """Serialización única y determinista de todo lo que escribe el conversor."""
    return json.dumps(datos, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def _escribir_destino(
    directorio: Path, definicion: DefLibro, huella: str, casos: list[CasoConvertido]
) -> list[Path]:
    directorio.mkdir(parents=True, exist_ok=True)
    escritos: list[Path] = []

    esperados = {f"{caso.caso_id}.json" for caso in casos} | {NOMBRE_INDICE}
    for sobrante in directorio.glob("*.json"):
        if sobrante.name not in esperados:
            sobrante.unlink()

    for caso in casos:
        ruta = directorio / f"{caso.caso_id}.json"
        ruta.write_text(_volcar(caso.a_json()), encoding="utf-8")
        escritos.append(ruta)

    indice = directorio / NOMBRE_INDICE
    indice.write_text(
        _volcar(
            {
                "casos": [caso.caso_id for caso in casos],
                "fase": definicion.fase,
                "libro": definicion.fichero,
                "sha256_libro": huella,
            }
        ),
        encoding="utf-8",
    )
    escritos.append(indice)
    return escritos


def convertir(
    dir_ground_truth: Path | str = RUTA_GROUND_TRUTH,
    dir_fixtures: Path | str = RUTA_FIXTURES,
) -> InformeConversion:
    """Convierte los seis libros a fixtures JSON. Todo o nada."""
    origen = Path(dir_ground_truth)
    destino = Path(dir_fixtures)
    if not origen.is_dir():
        raise ErrorConversion(
            f"no existe el directorio de ground truth '{origen}'. Los libros "
            f"no se versionan: pídeselos al humano o créalos desde la spec."
        )

    hallazgos: list[Hallazgo] = []
    convertidos: list[tuple[DefLibro, str, list[CasoConvertido]]] = []
    for definicion in LIBROS:
        huella, casos = _convertir_libro(definicion, origen / definicion.fichero, hallazgos)
        convertidos.append((definicion, huella, casos))

    if hallazgos:
        raise ErrorDatosSensibles(hallazgos)

    informe = InformeConversion()
    for definicion, huella, casos in convertidos:
        informe.libros.append(definicion.fichero)
        informe.casos_por_fase[definicion.fase] = len(casos)
        informe.ficheros.extend(
            _escribir_destino(destino / definicion.destino, definicion, huella, casos)
        )
    return informe


# --- CLI --------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    """`python -m evals.conversor`: convierte y resume; 1 si algo falla."""
    analizador = argparse.ArgumentParser(
        prog="python -m evals.conversor",
        description="Convierte los libros de ground truth en fixtures JSON.",
    )
    analizador.add_argument("--ground-truth", default=str(RUTA_GROUND_TRUTH))
    analizador.add_argument("--fixtures", default=str(RUTA_FIXTURES))
    opciones = analizador.parse_args(argv)

    try:
        informe = convertir(opciones.ground_truth, opciones.fixtures)
    except ErrorConversion as error:
        print(str(error), file=sys.stderr)
        return 1

    print(informe.resumen())
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
