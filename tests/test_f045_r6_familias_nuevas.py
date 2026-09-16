# tests/test_f045_r6_familias_nuevas.py
"""F-045 · R6: las familias que aún no existen, y sus pestañas nuevas.

Tres etiquetas del Excel apuntan a familias de documento que hoy NO están en
el catálogo: `combustible` (que existe, pero solo de línea), `grava` y
`ferreteria`. Ampliarlo es F-046 —`familias.py` es ruta sensible: su texto se
inyecta en el prompt de IA1—, así que esos casos **nacen ROJOS a propósito**.

Lo que aquí se comprueba es que el importador no los esconde ni los mezcla: se
escriben igual, en su pestaña, y el informe los agrupa bajo «la familia aún no
existe en el catálogo», aparte de los defectos reales de clasificación. Un
rojo esperado que parece regresión cuesta una tarde de investigación.
"""

from __future__ import annotations

import openpyxl
import pytest

from evals.revision import escritura, reparto, vocabulario
from evals.revision.modelos import FilaPlana

VOCAB = vocabulario.cargar()


def fila(numero, codigo, tipo, cif):
    return FilaPlana(
        numero_fila=numero,
        valores={
            "codigo_albaran": codigo, "tipo_albaran": tipo, "cif": cif,
            "origen_linea": "EN ALBARAN", "origen_contrato": "EN CONTRATO",
            "origen_precio": "DE CONTRATO", "origen_importe": "DE CONTRATO",
            "concepto": "lo que sea", "cantidad": 1, "unidad": "UD",
            "precio_unitario": 10, "importe": 10, "partida": "P1.01",
        },
    )


FILAS = [
    fila(2, "58826", "GRAVA", "B29679628"),
    fila(3, "2139643", "FERRETERIA", "A28733558"),
    fila(4, "J1000505", "GASOLEO", "B98976111"),
    fila(5, "H98634", "HORMIGON", "B04685541"),
]


def casos():
    casos = reparto.agrupar_por_albaran(FILAS, VOCAB)
    reparto.asignar_casos_id(casos, {})
    return casos


# --- El caso se escribe igual, y el informe lo agrupa aparte ---------------


def test_f045_r6_los_casos_de_familia_pendiente_se_escriben_igual():
    for caso in casos():
        tablas = reparto.repartir(caso, VOCAB)
        assert tablas["IA1"]["cabeceras"][0]["caso_id"] == caso.caso_id


def test_f045_r6_el_informe_agrupa_las_familias_que_aun_no_existen():
    pendientes = reparto.familias_pendientes(casos())
    assert pendientes == {
        "grava": ["GRA-001"],
        "ferreteria": ["FER-001"],
        "combustible": ["COM-001"],
    }


def test_f045_r6_una_familia_del_catalogo_no_se_agrupa_como_pendiente():
    assert "hormigon" not in reparto.familias_pendientes(casos())


# --- Las pestañas nuevas del banco ----------------------------------------


def test_f045_r6_el_conversor_declara_grava_y_ferreteria():
    from evals import conversor

    assert "Grava" in conversor.TIPOLOGIAS
    assert "Ferreteria" in conversor.TIPOLOGIAS


def test_f045_r6_las_pestanas_nuevas_copian_la_estructura_de_generico(tmp_path):
    libro = openpyxl.Workbook()
    plantilla = libro.active
    plantilla.title = "Generico-Suministros"
    plantilla.append(["TABLA 1 — CABECERAS"])
    plantilla.append(["caso_id", "fichero_albaran", "comentario"])
    ruta = tmp_path / "libro.xlsx"
    libro.save(ruta)
    libro.close()

    libro = openpyxl.load_workbook(ruta)
    creadas = escritura.asegurar_pestanas(libro, ("Grava", "Ferreteria"))
    assert creadas == ["Grava", "Ferreteria"]
    for pestana in ("Grava", "Ferreteria"):
        hoja = libro[pestana]
        assert hoja["A1"].value == "TABLA 1 — CABECERAS"
        assert [c.value for c in hoja[2]] == ["caso_id", "fichero_albaran", "comentario"]


def test_f045_r6_asegurar_pestanas_no_pisa_una_que_ya_existe(tmp_path):
    libro = openpyxl.Workbook()
    libro.active.title = "Generico-Suministros"
    libro.active.append(["TABLA 1"])
    hoja = libro.create_sheet("Grava")
    hoja.append(["NO ME TOQUES"])
    creadas = escritura.asegurar_pestanas(libro, ("Grava",))
    assert creadas == []
    assert libro["Grava"]["A1"].value == "NO ME TOQUES"


def test_f045_r6_sin_plantilla_no_se_inventa_una_pestana(tmp_path):
    libro = openpyxl.Workbook()
    libro.active.title = "CASOS"
    with pytest.raises(escritura.ErrorEscritura):
        escritura.asegurar_pestanas(libro, ("Grava",))
